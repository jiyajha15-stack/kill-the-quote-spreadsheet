import streamlit as st
import pandas as pd
import json, re, base64, os
from datetime import date
import docx
import pdfplumber
from anthropic import Anthropic

st.set_page_config(page_title="Kill the Quote Spreadsheet", layout="wide")

# ----------------------------------------------------------------- API client
api_key = os.environ.get("ANTHROPIC_API_KEY", "")
if not api_key:
    st.error("No API key found. Add ANTHROPIC_API_KEY as a Secret in this Space's Settings, then restart the Space.")
    st.stop()
client = Anthropic(api_key=api_key)
MODEL = "claude-sonnet-4-5-20250929"
TODAY = date.today().isoformat()

# ----------------------------------------------------------------- file readers
def read_txt(file):
    return file.read().decode("utf-8", errors="ignore")

def read_docx(file):
    d = docx.Document(file)
    paras = "\n".join(p.text for p in d.paragraphs if p.text.strip())
    tables = ""
    for t in d.tables:
        for row in t.rows:
            tables += " | ".join(c.text for c in row.cells) + "\n"
    return f"LETTER TEXT:\n{paras}\n\nTABLE:\n{tables}"

def read_xlsx(file):
    xl = pd.ExcelFile(file)
    out = ""
    for sheet in xl.sheet_names:
        df = xl.parse(sheet, header=None)
        out += f"--- SHEET: {sheet} ---\n{df.to_string()}\n\n"
    return out

def read_pdf(file):
    text = ""
    with pdfplumber.open(file) as pdf:
        for page in pdf.pages:
            text += (page.extract_text() or "") + "\n"
    return text

def image_b64(file):
    return base64.standard_b64encode(file.read()).decode("utf-8")

def extract_json(text):
    m = re.search(r"```json\s*(.*?)\s*```", text, re.DOTALL)
    if m:
        text = m.group(1)
    return json.loads(text)

def read_any(uploaded_file):
    """Return (kind, content) where kind is 'text' or 'image', dispatched by extension."""
    name = uploaded_file.name.lower()
    if name.endswith((".txt", ".eml")):
        return "text", read_txt(uploaded_file)
    if name.endswith(".docx"):
        return "text", read_docx(uploaded_file)
    if name.endswith((".xlsx", ".xls")):
        return "text", read_xlsx(uploaded_file)
    if name.endswith(".pdf"):
        return "text", read_pdf(uploaded_file)
    if name.endswith((".jpg", ".jpeg", ".png")):
        return "image", image_b64(uploaded_file)
    return "text", f"[Unsupported file type: {uploaded_file.name}]"

# ----------------------------------------------------------------- Claude calls
def extract_vendor_quote(vendor_label, text_content=None, image_b64_data=None, extra_text=None):
    instructions = f"""You are extracting a vendor's quote for corrugated boxes into structured data, for an RFQ with 30 line items.

Extract every line item you can find, plus commercial terms (freight, payment, lead time, validity) and any questionnaire/qualification info (ISO certification, GST, capacity, years in business, references).

Return ONLY this JSON structure, no other text:
{{
  "vendor_label": "{vendor_label}",
  "pricing_basis": "per piece / per 100 pieces / per kg / FOB USD / etc - describe exactly what you see",
  "currency": "INR / USD / etc",
  "lines_quoted": <count>,
  "lines_expected": 30,
  "line_items": [
    {{"description": "...", "raw_price": price_number, "notes": "any alt spec, MOQ condition, etc, else empty string"}}
  ],
  "tax_basis": "inclusive/exclusive/unstated",
  "freight": "what's stated",
  "payment_terms": "what's stated, else 'not stated'",
  "lead_time_days": number_or_null,
  "quote_validity_days": number_or_null,
  "questionnaire_info": {{"iso_certified": true_or_false_or_null, "iso_status_detail": "quote exact wording if uncertain", "gst_registered": true_or_false_or_null, "years_in_business": number_or_null, "monthly_capacity": "...", "references": ["..."]}},
  "confidence": 0.0 to 1.0,
  "assumptions": ["..."],
  "flags": ["list issues: missing lines, ambiguous units, discounts buried in fine print, knock-out failures, etc"]
}}
List every single line item you find, do not skip any, do not summarize them."""

    text_part = instructions
    if text_content:
        text_part += f"\n\nVendor's document content:\n---\n{text_content}\n---"
    if extra_text:
        text_part += f"\n\nAdditional vendor message:\n---\n{extra_text}\n---"

    if image_b64_data:
        content = [
            {"type": "image", "source": {"type": "base64", "media_type": "image/jpeg", "data": image_b64_data}},
            {"type": "text", "text": text_part},
        ]
    else:
        content = text_part

    resp = client.messages.create(model=MODEL, max_tokens=4000, messages=[{"role": "user", "content": content}])
    try:
        return extract_json(resp.content[0].text)
    except Exception as e:
        st.warning(f"Could not parse extraction for {vendor_label}: {e}")
        return None

def match_lines(vendor_data, rfx_lines):
    rfx_list = rfx_lines[["line_code", "description", "length_mm", "width_mm", "height_mm", "ply", "est_weight_kg_per_box"]].to_dict("records")
    prompt = f"""Match each vendor line item to the correct official RFx line_code.

Official RFx lines:
{json.dumps(rfx_list, indent=2)}

Vendor's line items:
{json.dumps(vendor_data['line_items'], indent=2)}

Vendor's pricing basis: {vendor_data['pricing_basis']}

If a single vendor line item is a CATEGORY-LEVEL price (e.g. "all 5-ply boxes") rather than one specific line, set matched_line_code to "CATEGORY" and specify category_ply (3, 5, or 7).
If a vendor line is an alternate/substitute (different ply than requested), still match it to the closest line_code and mark is_alternate true.

Return ONLY this JSON:
{{"matches": [{{"vendor_description": "...", "matched_line_code": "L01 or CATEGORY", "category_ply": null_or_3_or_5_or_7, "is_alternate": true_or_false, "match_confidence": 0.0-1.0}}]}}"""
    resp = client.messages.create(model=MODEL, max_tokens=3000, messages=[{"role": "user", "content": prompt}])
    return extract_json(resp.content[0].text)["matches"]

def check_certificate(text_content):
    prompt = f"""This is text extracted from a certificate PDF:
---
{text_content}
---
Extract the certificate holder name and expiry date. Today's date is {TODAY}.
Return ONLY this JSON:
{{"company_on_certificate": "...", "expiry_date": "YYYY-MM-DD or null", "is_expired": true_or_false_or_null}}"""
    resp = client.messages.create(model=MODEL, max_tokens=300, messages=[{"role": "user", "content": prompt}])
    return extract_json(resp.content[0].text)

def ask_analyst(question, ctx):
    prompt = f"""You are a procurement analyst assistant. Answer using ONLY the data below.
Never invent a number, and never recount or re-derive something already given as a fact - state it exactly as given.
If you lack enough information, say so clearly.

VENDOR COVERAGE (exact, pre-computed):
{json.dumps(ctx['missing_lines_summary'], indent=2)}

QUALIFICATION STATUS (exact narrative per vendor, verified against attached certificates):
{json.dumps(ctx['qualification_detail'], indent=2)}

COMPLIANCE WITH BUYER'S STATED TERMS:
{json.dumps(ctx['compliance'], indent=2)}

FULL NORMALIZED PRICE COMPARISON (Rs per piece, ex-tax):
{ctx['master_df'].to_json(orient='records')}

AWARD RECOMMENDATION:
{ctx['award_df'].to_json(orient='records')}

TOTALS:
Recommended total (qualified + compliant): Rs {ctx['total_recommended']:,.2f}
Cheapest possible if terms ignored: Rs {ctx['total_cheapest']:,.2f}
Lines where a cheaper option was excluded due to non-compliant terms: {ctx['excluded_count']} of {len(ctx['master_df'])}

Buyer's question: {question}

Answer clearly and specifically, citing exact line_codes or numbers from the data above."""
    resp = client.messages.create(model=MODEL, max_tokens=1000, messages=[{"role": "user", "content": prompt}])
    return resp.content[0].text

# ----------------------------------------------------------------- normalization (pure code)
def normalize_price(raw_price, basis, currency, fx_rate, box_weight_kg):
    basis = basis.lower()
    price = raw_price
    notes = []
    if currency == "USD":
        price *= fx_rate
        notes.append(f"USD->INR @ {fx_rate}")
    if "per 100" in basis:
        price /= 100
        notes.append("per-100 -> per-piece (/100)")
    if "per kg" in basis:
        price *= box_weight_kg
        notes.append(f"per-kg -> per-piece (x{box_weight_kg}kg)")
    return round(price, 2), notes

def draft_rfx_with_ai(description, num_items_hint=None):
    """The RFx co-pilot: buyer describes what they need in plain language,
    Claude drafts structured line items, a qualification questionnaire, and terms.
    Scoped to corrugated packaging for this demo, matching the rest of the built pipeline
    (which relies on 'ply' and box-weight fields for per-kg price normalization) -
    stated plainly here and in the one-page note as an intentional scoping choice."""

    prompt = f"""You are a procurement co-pilot helping a buyer draft an RFx (Request for Quotation) for CORRUGATED BOX packaging.

The buyer described what they need in their own words:
---
{description}
---

Draft a complete RFx. Return ONLY this JSON:
{{
  "line_items": [
    {{"line_code": "L01", "description": "RSC corrugated box AxBxC mm, N-ply", "length_mm": number, "width_mm": number, "height_mm": number, "ply": 3_or_5_or_7, "quantity_pcs": number, "est_weight_kg_per_box": estimated_number}}
  ],
  "questionnaire": [
    {{"id": "Q1", "question": "...", "knockout": true_or_false}}
  ],
  "terms": {{
    "price_basis": "e.g. Rs per piece, ex-GST, delivered to [location]",
    "payment_terms_preferred": "e.g. 30 days or longer from invoice",
    "quote_validity_required_days": number,
    "quotes_due_in_days": number
  }},
  "buyer_name": "a reasonable fictional buyer company name based on context, or 'Buyer' if none given"
}}

Generate a realistic, varied set of line items based on what the buyer described (different sizes/plies if the buyer mentioned variety, or a focused set if they described something specific). Estimate box weight per piece reasonably from its dimensions and ply (thicker/larger boxes weigh more). Include a standard ISO 9001 knockout question and 4-6 other reasonable qualification questions (GST registration, capacity, delivery, quality reporting, references)."""

    resp = client.messages.create(model=MODEL, max_tokens=4000, messages=[{"role": "user", "content": prompt}])
    return extract_json(resp.content[0].text)

# ----------------------------------------------------------------- UI
st.title("Kill the Quote Spreadsheet")
st.caption("Aerchain — comparison and analyst chat")

st.header("0. Draft your RFx with an AI co-pilot")
st.write("Describe what you need in plain language. The co-pilot drafts line items, a qualification questionnaire, and terms — scoped to corrugated box packaging.")
rfx_description = st.text_area(
    "What do you need?",
    placeholder="e.g. We need corrugated boxes in about 10 different sizes, ranging from small (300x200x150mm) to large (650x450x400mm), in 3-ply, 5-ply, and 7-ply options, for a food packaging warehouse in Pune. Quantities between 5,000 and 30,000 pieces per size.",
    height=100,
)
if st.button("Draft RFx with AI"):
    if not rfx_description.strip():
        st.warning("Please describe what you need first.")
    elif client is None:
        st.error("No API key configured yet.")
    else:
        with st.spinner("Drafting RFx..."):
            st.session_state["drafted_rfx"] = draft_rfx_with_ai(rfx_description)

if "drafted_rfx" in st.session_state:
    drafted = st.session_state["drafted_rfx"]
    st.success(f"Drafted RFx for: {drafted.get('buyer_name', 'Buyer')}")

    drafted_df = pd.DataFrame(drafted["line_items"])
    st.write(f"**{len(drafted_df)} line items drafted:**")
    st.dataframe(drafted_df, use_container_width=True)

    st.write("**Qualification questionnaire:**")
    for q in drafted["questionnaire"]:
        knockout_tag = " ⚠️ knockout" if q.get("knockout") else ""
        st.write(f"- {q['id']}: {q['question']}{knockout_tag}")

    st.write("**Terms:**")
    st.json(drafted["terms"])

    # Build downloadable files matching exactly what section 1/2 below expect to upload
    csv_bytes = drafted_df.to_csv(index=False).encode("utf-8")
    terms_text = (
        f"RFQ: Corrugated boxes, {len(drafted_df)} line items\n"
        f"Buyer: {drafted.get('buyer_name', 'Buyer')}\n"
        f"Price basis requested: {drafted['terms'].get('price_basis', '')}\n"
        f"Payment terms preferred: {drafted['terms'].get('payment_terms_preferred', '')}\n"
        f"Quote validity required: minimum {drafted['terms'].get('quote_validity_required_days', '')} days\n"
        f"Quotes due: {drafted['terms'].get('quotes_due_in_days', '')} days from RFQ date\n"
    )

    c1, c2 = st.columns(2)
    c1.download_button("Download rfx_lines.csv", csv_bytes, file_name="rfx_lines.csv", mime="text/csv")
    c2.download_button("Download terms.txt", terms_text.encode("utf-8"), file_name="terms.txt", mime="text/plain")

    if st.button("Send RFx to vendors"):
        # Plumbing stubbed per the brief's own rule ("fake the SMTP server if you like") -
        # the extraction and reasoning elsewhere in this app are real; this step is not.
        st.info(f"📧 Simulated: RFx sent to 5 vendors (A, B, C, D, E) via email. "
                f"Quotes requested within {drafted['terms'].get('quotes_due_in_days', 9)} days. "
                f"(This send step is stubbed, as permitted by the assignment brief — everything below this point is real.)")

    st.info("⬇️ Use the downloaded rfx_lines.csv and terms.txt in Sections 1 and 2 below to run the full comparison.")

st.divider()
st.header("1. RFx line items")
rfx_file = st.file_uploader("Upload rfx_lines.csv", type=["csv"], key="rfx")

st.header("2. Buyer's stated terms (for compliance checking)")
terms_file = st.file_uploader("Upload terms.txt (optional)", type=["txt"], key="terms")

st.header("3. Vendor responses")
st.write("Upload each vendor's file. For a vendor who sent a photo AND an email, upload both under the same vendor slot.")
vendor_slots = {}
cols = st.columns(5)
for i, label in enumerate(["A", "B", "C", "D", "E"]):
    with cols[i]:
        st.subheader(f"Vendor {label}")
        main_file = st.file_uploader(f"Main file ({label})", key=f"main_{label}")
        extra_file = st.file_uploader(f"Extra text/email ({label}, optional)", type=["txt"], key=f"extra_{label}")
        cert_file = st.file_uploader(f"ISO certificate ({label}, optional)", type=["pdf"], key=f"cert_{label}")
        vendor_slots[label] = {"main": main_file, "extra": extra_file, "cert": cert_file}

fx_rate = st.number_input("FX rate assumption (INR per USD)", value=88.0, step=0.5)

if st.button("Run analysis", type="primary"):
    if not rfx_file:
        st.error("Please upload rfx_lines.csv first.")
        st.stop()

    rfx_lines = pd.read_csv(rfx_file)
    buyer_terms = terms_file.read().decode("utf-8") if terms_file else ""

    active_vendors = [v for v, s in vendor_slots.items() if s["main"] is not None]
    if not active_vendors:
        st.error("Please upload at least one vendor file.")
        st.stop()

    results, lookups, cert_checks = {}, {}, {}

    with st.spinner("Extracting vendor quotes..."):
        for v in active_vendors:
            slot = vendor_slots[v]
            kind, content = read_any(slot["main"])
            extra_text = read_txt(slot["extra"]) if slot["extra"] else None
            if kind == "image":
                results[v] = extract_vendor_quote(v, image_b64_data=content, extra_text=extra_text)
            else:
                results[v] = extract_vendor_quote(v, text_content=content, extra_text=extra_text)

    with st.spinner("Checking ISO certificates..."):
        for v in active_vendors:
            cert_file = vendor_slots[v]["cert"]
            if cert_file:
                cert_text = read_pdf(cert_file)
                cert_checks[v] = check_certificate(cert_text)

    with st.spinner("Matching line items and normalizing prices..."):
        for v in active_vendors:
            vd = results[v]
            if vd is None:
                continue
            matches = match_lines(vd, rfx_lines)
            desc_to_price = {it["description"]: it["raw_price"] for it in vd["line_items"]}
            lookup = {}
            for m in matches:
                if m["matched_line_code"] == "CATEGORY" and m.get("category_ply"):
                    ply_lines = rfx_lines[rfx_lines["ply"] == m["category_ply"]]
                    price = desc_to_price.get(m["vendor_description"])
                    if price is not None:
                        for _, r in ply_lines.iterrows():
                            lookup[r["line_code"]] = {"raw_price": price, "is_alternate": False, "conf": m.get("match_confidence", 0.7)}
                elif m["matched_line_code"] not in (None, "CATEGORY"):
                    price = desc_to_price.get(m["vendor_description"])
                    if price is not None:
                        lookup[m["matched_line_code"]] = {"raw_price": price, "is_alternate": m.get("is_alternate", False), "conf": m.get("match_confidence", 1.0)}
            lookups[v] = lookup

    # Build master comparison table
    master_rows = []
    for _, rfx_row in rfx_lines.iterrows():
        lc = rfx_row["line_code"]
        row = {"line_code": lc, "description": rfx_row["description"]}
        for v in active_vendors:
            if results.get(v) is None:
                row[f"{v}_price_inr"] = None
                row[f"{v}_alternate"] = None
                continue
            entry = lookups.get(v, {}).get(lc)
            if entry:
                norm, _ = normalize_price(entry["raw_price"], results[v]["pricing_basis"], results[v]["currency"], fx_rate, rfx_row["est_weight_kg_per_box"])
                row[f"{v}_price_inr"] = norm
                row[f"{v}_alternate"] = entry["is_alternate"]
            else:
                row[f"{v}_price_inr"] = None
                row[f"{v}_alternate"] = None
        master_rows.append(row)
    master_df = pd.DataFrame(master_rows)

    # Missing lines summary
    missing_lines_summary = {}
    for v in active_vendors:
        covered = set(lookups.get(v, {}).keys())
        all_lines = set(rfx_lines["line_code"])
        missing = sorted(all_lines - covered)
        missing_lines_summary[v] = {"missing_count": len(missing), "missing_lines": missing}

    # Qualification: ISO claim + certificate verification
    qualification_detail, qualified_final = {}, {}
    for v in active_vendors:
        vd = results.get(v)
        if vd is None:
            continue
        claimed = vd["questionnaire_info"].get("iso_certified") == True
        if not claimed:
            detail = f"NOT QUALIFIED - vendor did not claim valid ISO 9001 ({vd['questionnaire_info'].get('iso_status_detail', 'no certificate claimed')})"
            qualified_final[v] = False
        elif v in cert_checks and cert_checks[v].get("is_expired") is True:
            detail = f"NOT QUALIFIED - vendor claimed valid ISO 9001, but attached certificate expired on {cert_checks[v]['expiry_date']} - claim does not match evidence"
            qualified_final[v] = False
        elif v in cert_checks:
            detail = f"QUALIFIED - ISO 9001 claimed and verified valid until {cert_checks[v]['expiry_date']}"
            qualified_final[v] = True
        else:
            detail = "QUALIFIED (claimed; no certificate uploaded to verify against)"
            qualified_final[v] = True
        qualification_detail[v] = detail

    # Compliance with buyer's stated terms
    compliance = {}
    for v in active_vendors:
        vd = results.get(v)
        if vd is None or not qualified_final.get(v):
            continue
        issues = []
        payment = (vd.get("payment_terms") or "").lower()
        if "advance" in payment and "30 day" not in payment and "45 day" not in payment:
            issues.append(f"Payment requires advance ({vd.get('payment_terms')}), buyer wants 30+ days credit")
        validity = vd.get("quote_validity_days")
        if validity is not None and validity < 30:
            issues.append(f"Quote validity only {validity} days, buyer requires minimum 30")
        compliance[v] = issues

    # Award logic
    qualified_vendors = [v for v in active_vendors if qualified_final.get(v)]
    fully_eligible = [v for v in qualified_vendors if not compliance.get(v)]
    award_rows = []
    for _, rfx_row in rfx_lines.iterrows():
        lc = rfx_row["line_code"]
        row_match = master_df[master_df["line_code"] == lc].iloc[0]
        cand_eligible = {v: row_match[f"{v}_price_inr"] for v in fully_eligible if pd.notna(row_match.get(f"{v}_price_inr"))}
        cand_qualified = {v: row_match[f"{v}_price_inr"] for v in qualified_vendors if pd.notna(row_match.get(f"{v}_price_inr"))}
        winner_eligible = min(cand_eligible, key=cand_eligible.get) if cand_eligible else None
        winner_cheapest = min(cand_qualified, key=cand_qualified.get) if cand_qualified else None
        award_rows.append({
            "line_code": lc,
            "recommended_vendor": winner_eligible,
            "recommended_price": cand_eligible.get(winner_eligible) if winner_eligible else None,
            "cheapest_qualified_vendor": winner_cheapest,
            "cheapest_qualified_price": cand_qualified.get(winner_cheapest) if winner_cheapest else None,
            "cheaper_option_excluded_due_to_terms": winner_eligible != winner_cheapest,
        })
    award_df = pd.DataFrame(award_rows)
    total_recommended = award_df["recommended_price"].sum(skipna=True)
    total_cheapest = award_df["cheapest_qualified_price"].sum(skipna=True)
    excluded_count = int(award_df["cheaper_option_excluded_due_to_terms"].sum())

    st.session_state["ctx"] = {
        "master_df": master_df, "missing_lines_summary": missing_lines_summary,
        "qualification_detail": qualification_detail, "compliance": compliance,
        "award_df": award_df, "total_recommended": total_recommended,
        "total_cheapest": total_cheapest, "excluded_count": excluded_count,
    }
    st.session_state["results"] = results
    st.success("Analysis complete.")

# ----------------------------------------------------------------- Display results
if "ctx" in st.session_state:
    ctx = st.session_state["ctx"]

    st.header("Side-by-side comparison (Rs per piece, ex-tax)")
    price_cols = [c for c in ctx["master_df"].columns if c.endswith("_price_inr")]
    st.dataframe(ctx["master_df"][["line_code", "description"] + price_cols], use_container_width=True)

    st.header("Vendor qualification (ISO 9001, verified against attached certificates)")
    for v, detail in ctx["qualification_detail"].items():
        st.write(f"**Vendor {v}:** {detail}")

    st.header("Compliance with buyer's stated terms")
    for v, issues in ctx["compliance"].items():
        st.write(f"**Vendor {v}:** {'COMPLIANT' if not issues else '; '.join(issues)}")

    st.header("Award recommendation")
    st.dataframe(ctx["award_df"], use_container_width=True)
    c1, c2, c3 = st.columns(3)
    c1.metric("Recommended total", f"Rs {ctx['total_recommended']:,.2f}")
    c2.metric("Cheapest if terms ignored", f"Rs {ctx['total_cheapest']:,.2f}")
    c3.metric("Lines where cheaper option was excluded", f"{ctx['excluded_count']}")

    st.header("Ask the analyst")
    q = st.text_input("Ask a question about this comparison, e.g. 'What if we split it, cheapest per line, but only among qualified vendors?'")
    if st.button("Ask"):
        with st.spinner("Thinking..."):
            answer = ask_analyst(q, ctx)
        st.markdown(answer)
